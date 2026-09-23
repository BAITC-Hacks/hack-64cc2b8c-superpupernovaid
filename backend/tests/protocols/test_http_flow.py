"""Full HTTP flow: real files/FFmpeg/SQL/renderers, deterministic model responses."""

import asyncio
import hashlib
import io
import shutil
import subprocess
from uuid import UUID

import pytest
from docx import Document
from fastapi.testclient import TestClient
from pypdf import PdfReader
from sqlalchemy import create_engine, event

from app.audio.ffmpeg import FfmpegAudioConverter
from app.audio.repository import AudioRepository
from app.audio.service import AudioPreprocessor
from app.bootstrap import get_audio_preprocessor, get_media_repository, get_media_service
from app.canonicalization.dependencies import (
    canonicalization_profile,
    get_canonicalization_service,
    get_canonicalization_speech_repository,
)
from app.canonicalization.models import CanonicalizationBatchResult, CanonicalizedSegmentOutput
from app.canonicalization.repository import CanonicalizationRepository
from app.canonicalization.service import TranscriptCanonicalizationService
from app.config import Settings, get_settings
from app.entrypoints.api import app
from app.infrastructure.database import Base
from app.intelligence.dependencies import get_analysis_coordinator
from app.media.probe import FFprobeMediaProbe
from app.media.repository import MediaRepository
from app.media.service import MediaService
from app.media.storage import LocalMediaStorage
from app.meetings.progress import ProcessingTracker
from app.meetings.repository import MeetingRepository, get_meetings
from app.protocols.dependencies import get_export_service, get_protocol_loader
from app.protocols.docx import DocxProtocolRenderer
from app.protocols.loader import ProtocolLoader
from app.protocols.repository import ExportRepository
from app.protocols.service import ProtocolExportService
from app.speech.alignment import SpeakerTranscriptAligner
from app.speech.dependencies import get_speech_audio_repository, get_speech_service, speech_profile
from app.speech.models import DiarizationResult, TranscriptionResult
from app.speech.repository import SpeechRepository
from app.speech.service import SpeechService
from app.tasks.repository import TaskRepository, get_tasks

from .test_analysis_integration import Runner, coordinator

pytestmark = pytest.mark.integration


class SpeechModel:
    def __init__(self):
        self.calls = []

    async def transcribe(self, audio):
        self.calls.append("asr")
        assert audio.sample_rate == 16000 and audio.channels == 1
        return TranscriptionResult(segments=[dict(start=0, end=1, text="Проверить отчёт.")])

    async def diarize(self, audio):
        self.calls.append("diarization")
        return DiarizationResult(segments=[dict(start=0, end=1, speaker_id="SPEAKER_00")])


class CanonicalModel:
    def __init__(self):
        self.calls = 0

    async def canonicalize_batch(self, batch):
        self.calls += 1
        return CanonicalizationBatchResult(
            segments=[
                CanonicalizedSegmentOutput(id=x.id, canonical_text=x.text, uncertain=False)
                for x in batch.targets
            ]
        )


@pytest.mark.parametrize("automatic", [False, True])
def test_upload_to_download_and_repeat(tmp_path, pdf_renderer, monkeypatch, automatic):
    if not shutil.which("ffmpeg") or not shutil.which("ffprobe"):
        pytest.skip("Requires real ffmpeg and ffprobe")
    engine = create_engine("sqlite:///" + str(tmp_path / "flow.sqlite"))

    @event.listens_for(engine, "connect")
    def foreign_keys(connection, _):
        connection.execute("PRAGMA foreign_keys=ON")

    Base.metadata.create_all(engine)
    settings = Settings(_env_file=None, meeting_tracing_enabled=False)
    storage = LocalMediaStorage(tmp_path / "storage")
    media_repo, audio_repo, speech_repo = (
        MediaRepository(engine),
        AudioRepository(engine),
        SpeechRepository(engine),
    )
    tracker = ProcessingTracker(engine)
    preprocessor = AudioPreprocessor(
        storage,
        FfmpegAudioConverter(FFprobeMediaProbe(10), 20),
        audio_repo,
        settings.audio_processing_config,
        tracker=tracker,
    )
    model, canonical_model, agents = SpeechModel(), CanonicalModel(), Runner()
    speech = SpeechService(
        model,
        model,
        SpeakerTranscriptAligner(),
        speech_repo,
        speech_profile(settings),
        tracker=tracker,
    )
    canonical = TranscriptCanonicalizationService(
        canonical_model,
        settings,
        CanonicalizationRepository(engine),
        canonicalization_profile(settings),
    )
    analysis = coordinator(engine, agents)
    exports = ProtocolExportService(
        {"docx": DocxProtocolRenderer(), "pdf": pdf_renderer},
        storage,
        ExportRepository(engine),
    )
    overrides = {
        get_settings: lambda: settings,
        get_meetings: lambda: MeetingRepository(engine),
        get_tasks: lambda: TaskRepository(engine),
        get_media_service: lambda: MediaService(
            storage,
            FFprobeMediaProbe(10),
            media_repo,
            1024 * 1024,
            frozenset({"wav"}),
        ),
        get_media_repository: lambda: media_repo,
        get_audio_preprocessor: lambda: preprocessor,
        get_speech_audio_repository: lambda: audio_repo,
        get_speech_service: lambda: speech,
        get_canonicalization_speech_repository: lambda: speech_repo,
        get_canonicalization_service: lambda: canonical,
        get_analysis_coordinator: lambda: analysis,
        get_protocol_loader: lambda: ProtocolLoader(engine),
        get_export_service: lambda: exports,
    }
    from app.processing import dependencies as processing_di
    from app.processing.repository import ProcessingRepository
    from app.processing.service import MeetingProcessingService, SubmitMeetingProcessing

    run_repo = ProcessingRepository(engine)
    deliveries = []

    class Queue:
        def enqueue(self, ident):
            deliveries.append(ident)

    for factory in (
        get_media_repository,
        get_audio_preprocessor,
        get_speech_service,
        get_canonicalization_service,
        get_protocol_loader,
        get_export_service,
    ):
        monkeypatch.setattr(processing_di, factory.__name__, overrides[factory])
    monkeypatch.setattr(processing_di, "analysis_coordinator", lambda: analysis)
    monkeypatch.setattr(processing_di, "preflight", lambda *_: None)
    overrides[processing_di.get_processing_repository] = lambda: run_repo
    overrides[processing_di.get_submit_meeting_processing] = lambda: SubmitMeetingProcessing(
        run_repo, Queue(), lambda *_: None
    )
    source = tmp_path / "meeting.wav"
    subprocess.run(
        [
            "ffmpeg",
            "-v",
            "error",
            "-f",
            "lavfi",
            "-i",
            "sine=duration=2:sample_rate=48000",
            "-ac",
            "2",
            str(source),
        ],
        check=True,
        capture_output=True,
        timeout=20,
    )
    source_hash = hashlib.sha256(source.read_bytes()).hexdigest()
    app.dependency_overrides.update(overrides)
    try:
        with TestClient(app) as client:
            created = client.post(
                "/api/v1/meetings",
                json={
                    "title": "Контрольный прогон",
                    "recording_consent_confirmed": True,
                },
            )
            assert created.status_code == 201, created.text
            base = "/api/v1/meetings/" + created.json()["id"]
            assert client.post(base + "/analyze").status_code == 409
            with source.open("rb") as stream:
                uploaded = client.post(
                    base + "/media", files={"file": (source.name, stream, "audio/wav")}
                )
            assert uploaded.status_code == 201, uploaded.text
            media = uploaded.json()
            audio_url = base + "/media/" + media["id"]
            if automatic:
                queued = client.post(base + "/process", json={"media_id": media["id"]})
                assert queued.status_code == 202, queued.text
                assert queued.json()["status"] == "queued" and len(deliveries) == 1
                workflow = MeetingProcessingService(run_repo, processing_di.ExistingMeetingStages())
                asyncio.run(workflow.execute(deliveries[0]))
                finished = client.get(base + "/processing-run").json()
                assert finished["status"] == "completed", finished
                assert set(finished["steps"].values()) == {"completed"}
                assert len(finished["exports"]) == 2
                assert client.get(base + "/processing").json()["status"] == "ready"
                again = client.post(base + "/process", json={"media_id": media["id"]})
                assert again.json()["id"] == queued.json()["id"]
                assert len(deliveries) == 1
            audio = client.post(audio_url + "/preprocess")
            assert audio.status_code == 200, audio.text
            assert audio.json()["sample_rate"] == 16000
            assert audio.json()["channels"] == 1
            transcript = client.post(audio_url + "/speech")
            assert transcript.status_code == 200, transcript.text
            if not automatic:
                assert client.post(base + "/analyze").status_code == 409
            clean = client.post(audio_url + "/canonicalize")
            assert clean.status_code == 200, clean.text
            assert clean.json()["segments"][0]["id"] == transcript.json()["segments"][0]["id"]
            assert (
                client.patch(
                    base + "/participants/SPEAKER_00", json={"display_name": "Әлия"}
                ).status_code
                == 200
            )
            result = client.post(base + "/analyze")
            assert result.status_code == 200, result.text
            assert result.json() == client.get(base + "/result").json()
            public = client.get(base + "/transcript").json()["items"]
            refs = result.json()["tasks"][0]["source_segment_ids"]
            assert UUID(refs[0]) and refs[0] == public[0]["id"]
            assert client.get(base + "/processing").json()["status"] == "ready"
            assert client.get(base + "/tasks").json()["items"] == result.json()["tasks"]
            for format in ("docx", "pdf"):
                generated = client.post(base + "/exports", json={"format": format})
                assert generated.status_code == 200, generated.text
                download = client.get(base + "/exports/" + format)
                assert download.status_code == 200, download.text
                assert hashlib.sha256(download.content).hexdigest() == generated.json()["sha256"]
                if format == "pdf":
                    text = "\n".join(
                        p.extract_text() for p in PdfReader(io.BytesIO(download.content)).pages
                    )
                else:
                    doc = Document(io.BytesIO(download.content))
                    text = "\n".join(p.text for p in doc.paragraphs)
                    text += "\n".join(c.text for t in doc.tables for r in t.rows for c in r.cells)
                assert "Контрольный прогон" in text and "Проверить" in text and "Әлия" in text
                assert (
                    client.post(base + "/exports", json={"format": format}).json()["id"]
                    == generated.json()["id"]
                )
            assert client.post(audio_url + "/preprocess").json()["id"] == audio.json()["id"]
            assert client.post(audio_url + "/speech").json() == transcript.json()
            assert client.post(audio_url + "/canonicalize").json() == clean.json()
            repeated = client.post(base + "/analyze")
            assert repeated.status_code == 200, repeated.text
            for field in ("tasks", "decisions", "summary", "analysis_details"):
                assert repeated.json()[field] == result.json()[field]
            assert repeated.json()["meeting"]["processing_status"] == "ready"
            assert model.calls == ["asr", "diarization"]
            assert canonical_model.calls == 1
            assert agents.calls == ["extraction", "resolver", "summary", "review"]
            with storage.open(media["storage_key"]) as original:
                assert hashlib.file_digest(original, "sha256").hexdigest() == source_hash
    finally:
        for key in overrides:
            app.dependency_overrides.pop(key, None)
        engine.dispose()
