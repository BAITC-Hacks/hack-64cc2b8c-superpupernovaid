import asyncio
from pathlib import Path

import pytest
from docx import Document
from pypdf import PdfReader

from app.protocols.docx import DocxProtocolRenderer
from app.protocols.errors import ProtocolDocumentRenderError
from app.protocols.models import ProtocolAction
from app.protocols.pdf import PdfProtocolRenderer
from app.protocols.presentation import timestamp


def docx_text(path):
    doc = Document(path)
    return "\n".join(
        [p.text for p in doc.paragraphs]
        + [c.text for t in doc.tables for row in t.rows for c in row.cells]
    )


def test_docx_pdf_same_content_and_no_mutation(protocol, pdf_renderer, tmp_path):
    before = protocol.model_dump_json()
    docx_path = tmp_path / "test.docx"
    pdf_path = tmp_path / "test.pdf"
    asyncio.run(DocxProtocolRenderer().render(protocol, docx_path))
    asyncio.run(pdf_renderer.render(protocol, pdf_path))
    assert pdf_path.read_bytes().startswith(b"%PDF-")
    docx_content = docx_text(docx_path)
    pdf_content = "\n".join(page.extract_text() for page in PdfReader(pdf_path).pages)
    for content in (docx_content, pdf_content):
        for expected in [
            "ПРОТОКОЛ СОВЕЩАНИЯ",
            protocol.summary,
            "Данияр Серикович",
            "Әлия Қанатқызы",
            "Подготовить претензию поставщику.",
            "До конца недели",
            "Не определён",
            "Не указан",
            "[01:34:27] SPEAKER_03",
            "Ертең отчетты до обеда жіберіңіз.",
            "Обсудим поставку сырья.",
            "цена < 100 & качество > 90.",
        ]:
            assert " ".join(expected.split()) in " ".join(content.split())
        assert "seg_001" not in content and "Обсудим сырьё." not in content
    assert protocol.model_dump_json() == before
    original_path = tmp_path / "original.docx"
    asyncio.run(DocxProtocolRenderer("original").render(protocol, original_path))
    assert "Обсудим сырьё." in docx_text(original_path)
    assert "Обсудим поставку сырья." not in docx_text(original_path)


def test_empty_sections_and_absent_status(protocol, pdf_renderer, tmp_path):
    empty = protocol.model_copy(
        update={
            "summary": None,
            "topics": (),
            "decisions": (),
            "participants": (),
            "action_items": (ProtocolAction(text="Only task"),),
            "transcript": (),
        }
    )
    for renderer, suffix in [(DocxProtocolRenderer(), "docx"), (pdf_renderer, "pdf")]:
        path = tmp_path / ("empty." + suffix)
        asyncio.run(renderer.render(empty, path))
        text = (
            docx_text(path)
            if suffix == "docx"
            else "\n".join(p.extract_text() for p in PdfReader(path).pages)
        )
        for absent in [
            "КРАТКОЕ СОДЕРЖАНИЕ",
            "ОСНОВНЫЕ ТЕМЫ",
            "ПРИНЯТЫЕ РЕШЕНИЯ",
            "СТЕНОГРАММА",
            "Статус",
        ]:
            assert absent not in text
        assert "Only task" in text


def test_long_task_splits_across_pdf_pages(protocol, pdf_renderer, tmp_path):
    text = ("Очень длинное поручение Қ Ғ Ә Ө Ұ Ү І Ң. " * 450) + "FINAL_TOKEN"
    large = protocol.model_copy(
        update={"action_items": (ProtocolAction(text=text, assignee="Длинное имя " * 10),)}
    )
    path = tmp_path / "long.pdf"
    asyncio.run(pdf_renderer.render(large, path))
    reader = PdfReader(path)
    assert len(reader.pages) > 3
    assert "FINAL_TOKEN" in "\n".join(p.extract_text() for p in reader.pages)


def test_font_failure_is_controlled(protocol, tmp_path):
    renderer = PdfProtocolRenderer(Path("/missing/secret/font.ttf"), Path("/missing/bold.ttf"))
    with pytest.raises(ProtocolDocumentRenderError) as error:
        asyncio.run(renderer.render(protocol, tmp_path / "out.pdf"))
    assert "secret" not in str(error.value)


@pytest.mark.parametrize(
    "seconds,expected",
    [(125.4, "00:02:05"), (5667.9, "01:34:27"), (360000, "100:00:00"), (59.999, "00:00:59")],
)
def test_timestamps(seconds, expected):
    assert timestamp(seconds) == expected
