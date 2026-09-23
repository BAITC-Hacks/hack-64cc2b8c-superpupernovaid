import asyncio
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

from app.protocols.errors import ProtocolDocumentRenderError
from app.protocols.models import MeetingProtocol, TextMode
from app.protocols.presentation import action_table, segment_heading, segment_text


class DocxProtocolRenderer:
    def __init__(self, mode: TextMode = "canonical"):
        self.mode = mode
        self.revision = "docx-v1-" + mode

    async def render(self, protocol: MeetingProtocol, destination: Path) -> Path:
        try:
            await asyncio.to_thread(self._render, protocol, destination)
            return destination
        except Exception:
            raise ProtocolDocumentRenderError from None

    def _render(self, protocol, destination):
        doc = Document()
        # python-docx's template may inherit decorative paragraph borders.
        for style in doc.styles:
            for border in style.element.xpath(".//w:pBdr"):
                border.getparent().remove(border)
        doc.core_properties.title = protocol.title
        doc.core_properties.author = "Qoryt"
        section = doc.sections[0]
        section.page_width, section.page_height = Cm(21), Cm(29.7)
        section.top_margin = section.bottom_margin = Cm(2)
        section.left_margin = section.right_margin = Cm(2)
        for name, size in [("Normal", 11), ("Title", 22), ("Heading 1", 15), ("Heading 2", 11)]:
            style = doc.styles[name]
            style.font.name = "DejaVu Sans"
            style.font.size = Pt(size)
            style.font.color.rgb = RGBColor(0, 0, 0)
            style.paragraph_format.space_after = Pt(6)
            if name in ("Heading 1", "Heading 2"):
                style.paragraph_format.keep_with_next = True
        doc.styles["Normal"].paragraph_format.line_spacing = 1.12
        doc.add_paragraph("ПРОТОКОЛ СОВЕЩАНИЯ", "Title")
        doc.add_paragraph("Название: " + protocol.title)
        if protocol.scheduled_at:
            doc.add_paragraph("Дата: " + protocol.scheduled_at.isoformat())
        if protocol.participants:
            doc.add_paragraph("Участники", "Heading 1")
            for person in protocol.participants:
                doc.add_paragraph(person.participant_name or person.speaker_id, "List Bullet")
        if protocol.summary:
            doc.add_paragraph("КРАТКОЕ СОДЕРЖАНИЕ", "Heading 1")
            doc.add_paragraph(protocol.summary)
        for title, values in [
            ("ОСНОВНЫЕ ТЕМЫ", protocol.topics),
            ("ПРИНЯТЫЕ РЕШЕНИЯ", tuple(x.text for x in protocol.decisions)),
        ]:
            if values:
                doc.add_paragraph(title, "Heading 1")
                for number, value in enumerate(values, 1):
                    doc.add_paragraph(f"{number}. {value}")
        if protocol.action_items:
            doc.add_paragraph("ПОРУЧЕНИЯ", "Heading 1")
            header, rows = action_table(protocol)
            table = doc.add_table(rows=1, cols=len(header))
            table.style = "Table Grid"
            table.alignment = WD_TABLE_ALIGNMENT.CENTER
            table.autofit = False
            widths = [0.8, 6.2, 3.4, 3.3, 3.3] if len(header) == 5 else [0.8, 8, 4, 4.2]
            for col, width in zip(table.columns, widths):
                col.width = Cm(width)
            for cell, text, width in zip(table.rows[0].cells, header, widths):
                cell.width = Cm(width)
                cell.text = text
                for run in cell.paragraphs[0].runs:
                    run.bold = True
            repeat = OxmlElement("w:tblHeader")
            table.rows[0]._tr.get_or_add_trPr().append(repeat)
            for values in rows:
                cells = table.add_row().cells
                for cell, text, width in zip(cells, values, widths):
                    cell.width = Cm(width)
                    cell.text = text
            for index, row in enumerate(table.rows):
                for cell in row.cells:
                    props = cell._tc.get_or_add_tcPr()
                    shade = OxmlElement("w:shd")
                    shade.set(qn("w:fill"), "EDEFF2" if index == 0 else "FFFFFF")
                    props.append(shade)
                    borders = OxmlElement("w:tcBorders")
                    for edge in ("top", "left", "bottom", "right"):
                        el = OxmlElement("w:" + edge)
                        el.set(qn("w:val"), "single")
                        el.set(qn("w:sz"), "4")
                        el.set(qn("w:color"), "D9D9D9")
                        borders.append(el)
                    props.append(borders)
                    for paragraph in cell.paragraphs:
                        paragraph.paragraph_format.space_after = Pt(4)
                        for run in paragraph.runs:
                            run.font.size = Pt(9)
        if protocol.transcript:
            heading = doc.add_paragraph("СТЕНОГРАММА", "Heading 1")
            heading.paragraph_format.page_break_before = True
            for seg in protocol.transcript:
                doc.add_paragraph(segment_heading(seg), "Heading 2")
                doc.add_paragraph(segment_text(seg, self.mode))
        doc.save(destination)
